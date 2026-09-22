"""
PPTX Generator — Creates PowerPoint presentations from slide data.
Uses python-pptx with a blue theme matching the ReviewMind website.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os


# ─── Theme colors (matching ReviewMind website) ───
BLUE_PRIMARY = RGBColor(0x25, 0x63, 0xEB)      # blue-600
BLUE_DARK = RGBColor(0x1E, 0x40, 0xAF)         # blue-800
BLUE_BG = RGBColor(0xEF, 0xF6, 0xFF)           # blue-50
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_DARK = RGBColor(0x1F, 0x29, 0x37)         # gray-800
GRAY_TEXT = RGBColor(0x4B, 0x55, 0x63)         # gray-600


def _add_background(slide, color=WHITE):
    """Fill slide background with a solid color."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_top_bar(slide, prs, height=Inches(0.15), color=BLUE_PRIMARY):
    """Add a thin colored bar at the top of the slide."""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, height
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    return bar


def _add_footer(slide, prs, text="ReviewMind"):
    """Add a small footer text at the bottom."""
    footer = slide.shapes.add_textbox(
        Inches(0.5), prs.slide_height - Inches(0.5),
        prs.slide_width - Inches(1), Inches(0.35)
    )
    tf = footer.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    for run in p.runs:
        run.font.size = Pt(10)
        run.font.color.rgb = GRAY_TEXT
        run.font.name = 'Calibri'


def _add_title_slide(prs, slide_data):
    """Create the title slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _add_background(slide, BLUE_BG)

    # Left accent bar
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.3), prs.slide_height
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE_PRIMARY
    bar.line.fill.background()

    # Title
    title_box = slide.shapes.add_textbox(
        Inches(1), Inches(2.2), prs.slide_width - Inches(2), Inches(2)
    )
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.text = slide_data.get('title', 'Untitled Presentation')
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    for run in p.runs:
        run.font.size = Pt(44)
        run.font.bold = True
        run.font.color.rgb = BLUE_DARK
        run.font.name = 'Calibri'

    # Subtitle
    if slide_data.get('subtitle'):
        sub_box = slide.shapes.add_textbox(
            Inches(1), Inches(4.3), prs.slide_width - Inches(2), Inches(1)
        )
        tf2 = sub_box.text_frame
        tf2.word_wrap = True
        tf2.text = slide_data['subtitle']
        p2 = tf2.paragraphs[0]
        for run in p2.runs:
            run.font.size = Pt(20)
            run.font.color.rgb = GRAY_TEXT
            run.font.name = 'Calibri'

    # Top bar
    _add_top_bar(slide, prs, height=Inches(0.15), color=BLUE_PRIMARY)
    # Move top bar to bottom
    slide.shapes[-1].top = prs.slide_height - Inches(0.15)


def _add_content_slide(prs, slide_data):
    """Create a standard content slide with title + bullet points."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_background(slide, WHITE)
    _add_top_bar(slide, prs)
    _add_footer(slide, prs)

    # Title
    title_box = slide.shapes.add_textbox(
        Inches(0.6), Inches(0.4), prs.slide_width - Inches(1.2), Inches(0.9)
    )
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.text = slide_data.get('title', '')
    p = tf.paragraphs[0]
    for run in p.runs:
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = BLUE_DARK
        run.font.name = 'Calibri'

    # Underline accent
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.6), Inches(1.25), Inches(1.5), Inches(0.05)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE_PRIMARY
    line.line.fill.background()

    # Bullet content
    content_box = slide.shapes.add_textbox(
        Inches(0.6), Inches(1.5),
        prs.slide_width - Inches(1.2), prs.slide_height - Inches(2.2)
    )
    tf_content = content_box.text_frame
    tf_content.word_wrap = True

    bullets = slide_data.get('bullets', [])
    if not bullets:
        bullets = [slide_data.get('content', '')]

    for i, bullet in enumerate(bullets):
        if i == 0:
            para = tf_content.paragraphs[0]
        else:
            para = tf_content.add_paragraph()
        para.text = f"• {bullet}"
        para.space_after = Pt(10)
        para.alignment = PP_ALIGN.LEFT
        for run in para.runs:
            run.font.size = Pt(18)
            run.font.color.rgb = GRAY_DARK
            run.font.name = 'Calibri'


def generate_pptx(slides_data, output_path):
    """
    Generate a PPTX file from slide data.

    Args:
        slides_data (dict): {
            'title': str,
            'subtitle': str (optional),
            'slides': [ {'title': str, 'bullets': [str, ...]}, ... ]
        }
        output_path (str): Where to save the .pptx file.

    Returns:
        str: Path to the generated file.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    # Title slide
    _add_title_slide(prs, slides_data)

    # Content slides
    for slide_data in slides_data.get('slides', []):
        _add_content_slide(prs, slide_data)

    # Ensure output dir exists
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    return output_path